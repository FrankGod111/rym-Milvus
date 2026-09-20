"""Document text extraction helpers for uploaded ERP files."""

from __future__ import annotations

import csv
import io
import json
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Any


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".json",
    ".log",
    ".xml",
    ".html",
    ".htm",
    ".yaml",
    ".yml",
}


class ParseResult(dict):
    text: str
    parser: str
    status: str
    error: str


class _HTMLTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        text = data.strip()
        if text:
            self.parts.append(text)

    def text(self) -> str:
        return "\n".join(self.parts)


def _decode_bytes(raw: bytes) -> str:
    for encoding in ["utf-8", "utf-8-sig", "gb18030", "latin-1"]:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _parse_csv(text: str) -> str:
    reader = csv.reader(io.StringIO(text))
    rows = ["\t".join(row) for row in reader]
    return "\n".join(rows)


def _parse_json(text: str) -> str:
    try:
        return json.dumps(json.loads(text), ensure_ascii=False, indent=2)
    except json.JSONDecodeError:
        return text


def _parse_html(text: str) -> str:
    parser = _HTMLTextParser()
    parser.feed(text)
    parsed = parser.text()
    return parsed or re.sub(r"<[^>]+>", " ", text)


def _parse_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF parsing requires pypdf. Install requirements.txt first.") from exc
    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()


def _parse_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX parsing requires python-docx. Install requirements.txt first.") from exc
    doc = Document(str(path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    table_rows: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            table_rows.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join([*paragraphs, *table_rows]).strip()


def _parse_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("XLSX parsing requires openpyxl. Install requirements.txt first.") from exc
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                parts.append("\t".join(values))
    return "\n".join(parts).strip()


def extract_text(path: Path, mime_type: str = "", file_name: str = "") -> dict[str, Any]:
    suffix = Path(file_name or path.name).suffix.lower()
    try:
        if suffix == ".pdf" or mime_type == "application/pdf":
            text = _parse_pdf(path)
            parser = "pypdf"
        elif suffix == ".docx":
            text = _parse_docx(path)
            parser = "python-docx"
        elif suffix in {".xlsx", ".xlsm"}:
            text = _parse_xlsx(path)
            parser = "openpyxl"
        elif suffix in TEXT_EXTENSIONS or mime_type.startswith("text/"):
            raw_text = _decode_bytes(path.read_bytes())
            if suffix == ".csv":
                text = _parse_csv(raw_text)
                parser = "csv"
            elif suffix == ".json":
                text = _parse_json(raw_text)
                parser = "json"
            elif suffix in {".html", ".htm"}:
                text = _parse_html(raw_text)
                parser = "html"
            else:
                text = raw_text
                parser = "text"
        else:
            return {
                "text": "",
                "parser": "unsupported",
                "status": "unsupported",
                "error": f"Unsupported file type: {suffix or mime_type or 'unknown'}",
            }
        return {
            "text": text,
            "parser": parser,
            "status": "parsed" if text.strip() else "empty",
            "error": "" if text.strip() else "No text extracted from file.",
        }
    except Exception as exc:  # keep upload durable even when parsing fails
        return {"text": "", "parser": "error", "status": "failed", "error": str(exc)}
